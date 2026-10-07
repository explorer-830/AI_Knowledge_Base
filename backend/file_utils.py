from pathlib import Path
import logging
from docx import Document
from openpyxl import load_workbook
from backend.pdf_utils import read_pdf

logger = logging.getLogger(__name__)
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".txt"}


def read_word(file_path: str):
    document = Document(file_path)
    parts = []
    for block in document.iter_inner_content():
        if hasattr(block, "text"):
            if block.text.strip():
                parts.append(block.text)
        else:
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_excel(file_path: str):
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    parts = []
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [f"第{column}列: {value}"
                         for column, value in enumerate(row, 1)
                         if value is not None and str(value).strip()]
                if cells:
                    parts.append(f"工作表: {sheet.title}, 第{row_number}行 | "
                                 + " | ".join(cells))
    finally:
        workbook.close()
    return "\n".join(parts)


def read_txt(file_path: str):
    content = Path(file_path).read_bytes()
    encodings = ("utf-16",) if content.startswith((b"\xff\xfe", b"\xfe\xff")) \
        else ("utf-8-sig", "gb18030")
    for encoding in encodings:
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" in text:
            raise ValueError("TXT 文件包含非文本内容")
        return text
    raise ValueError("TXT 编码无法识别，请转为 UTF-8 后上传")


def read_document(file_path: str):
    suffix = Path(file_path).suffix.lower()
    readers = {".pdf": read_pdf, ".docx": read_word,
               ".xlsx": read_excel, ".txt": read_txt}
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("仅支持 PDF、DOCX、XLSX、TXT 文件")
    try:
        text = readers[suffix](file_path)
    except (ValueError, OSError):
        raise
    except Exception as exc:
        logger.warning("文档解析失败: %s", suffix, exc_info=True)
        raise ValueError("文件无法解析，请检查文件是否损坏、加密或格式不匹配") from exc
    if not text.strip():
        raise ValueError("文件没有可提取的文本，请上传包含文字或单元格内容的文件")
    return text
